"""Build a Vietnamese PDF/Word report from measured artifacts and reviewed source."""

from __future__ import annotations

import ast
import hashlib
import json
import math
from pathlib import Path
import re
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
import pandas as pd
import pymupdf
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    Flowable,
    Image,
    PageBreak,
    Paragraph,
    Preformatted,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents

ROOT = Path(__file__).resolve().parent
ART = ROOT / "artifacts"
OUT = ROOT / "deliverables" / "tieng_viet"
OUT.mkdir(parents=True, exist_ok=True)
STEM = "DDM501_Assignment2_25MS13312_Nguyễn_Minh_Trường"
FONTDIR = Path("C:/Windows/Fonts")
for family, prefix in [("TimesVN", "times"), ("ArialVN", "arial")]:
    for suffix, ending in [("", ""), ("-Bold", "bd"), ("-Italic", "i"), ("-BoldItalic", "bi")]:
        pdfmetrics.registerFont(TTFont(family + suffix, str(FONTDIR / (prefix + ending + ".ttf"))))
    pdfmetrics.registerFontFamily(
        family,
        normal=family,
        bold=family + "-Bold",
        italic=family + "-Italic",
        boldItalic=family + "-BoldItalic",
    )
pdfmetrics.registerFont(TTFont("CodeVN", str(FONTDIR / "consola.ttf")))
WIDTH = A4[0] - 108
styles = getSampleStyleSheet()
styles.add(
    ParagraphStyle(
        "BodyVN",
        fontName="TimesVN",
        fontSize=12,
        leading=17,
        alignment=TA_JUSTIFY,
        spaceAfter=8,
        textColor=colors.HexColor("#243247"),
    )
)
styles.add(ParagraphStyle("CellVN", fontName="TimesVN", fontSize=9.5, leading=12.5))
styles.add(
    ParagraphStyle(
        "CaptionVN",
        fontName="TimesVN-Italic",
        fontSize=9.5,
        leading=13,
        spaceAfter=9,
        textColor=colors.HexColor("#466077"),
    )
)
styles.add(
    ParagraphStyle(
        "CodeVN",
        fontName="CodeVN",
        fontSize=7.1,
        leading=9.3,
        spaceBefore=5,
        spaceAfter=10,
        backColor=colors.HexColor("#f3f5f7"),
    )
)
styles.add(
    ParagraphStyle(
        "Cover",
        fontName="TimesVN-Bold",
        fontSize=23,
        leading=31,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#123e58"),
        spaceAfter=20,
    )
)
styles.add(
    ParagraphStyle(
        "Meta", fontName="TimesVN", fontSize=12, leading=19, alignment=TA_CENTER, spaceAfter=12
    )
)
for name, size, leading in [("Heading1", 17, 23), ("Heading2", 12.5, 18), ("Heading3", 11, 16)]:
    styles[name].fontName = "TimesVN-Bold"
    styles[name].fontSize = size
    styles[name].leading = leading
    styles[name].textColor = colors.HexColor("#123e58")
    styles[name].spaceBefore = 12
    styles[name].spaceAfter = 8


def read_json(name: str) -> dict:
    """Read completed experiment evidence; missing artifacts fail the report build."""
    return json.loads((ART / name).read_text(encoding="utf-8"))


def markdown_table(headers: list[str], rows: list[list[str]]) -> str:
    """Render a simple Markdown table that PDF and Word exporters share."""
    return "\n".join(
        ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
        + ["| " + " | ".join(row) + " |" for row in rows]
    )


def function_source(path: str, name: str) -> str:
    """Extract actual function text, maintaining consistency with executable code."""
    source = (ROOT / path).read_text(encoding="utf-8")
    tree = ast.parse(source)
    node = next(
        n
        for n in ast.walk(tree)
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name == name
    )
    return "\n".join(source.splitlines()[node.lineno - 1 : node.end_lineno])


def generated_sections() -> dict[str, str]:
    """Turn measured evidence into report tables and cautious decision analysis."""
    data = read_json("data_manifest.json")
    selection = read_json("selection.json")
    test = read_json("test_results.json")
    registry = read_json("registry.json")
    environment = read_json("environment.json")
    results = pd.read_csv(ART / "experiment_results.csv").set_index("id")
    chosen, baseline = selection["selected_id"], selection["baseline_id"]
    specs = json.loads((ART / "experiment_matrix.json").read_text(encoding="utf-8"))
    replacements = {}
    replacements["data_summary"] = (
        f"Thực nghiệm dùng {data['series']} chuỗi (4 cửa hàng × 12 SKU), "
        f"{data['raw_rows']:,} dòng nguồn, từ {data['date_min']} đến {data['date_max']}; "
        f"{data['example_rows']:,} mẫu origin/horizon trước label mask. Seed=501; "
        f"tỷ lệ ngày available thực sinh là {100 * data['availability_fraction']:.2f}%. "
        "Phạm vi này thu nhỏ hệ thống giả định, không thay đổi quy mô production của Assignment 1."
    )
    descriptions = {
        "E01": ("Mean 28", "Trung bình 28 ngày", "Baseline chính sách hiện tại"),
        "E02": ("Seasonal 4 tuần", "Cùng thứ 4 tuần", "Baseline mùa vụ"),
        "E03": ("Seasonal + promo", "Hệ số promo học trên train", "Baseline mạnh đã hiệu chỉnh"),
        "E04": ("Ridge", "alpha=10; full", "Đối chiếu mô hình đơn giản"),
        "E05": ("LGBM Poisson", "31 lá; lr=0,05; full", "Ứng viên A1 cụ thể hóa"),
        "E06": ("LGBM Poisson", "15 lá; lr=0,05; full", "Giảm capacity so E05"),
        "E07": ("LGBM Poisson", "63 lá; lr=0,05; full", "Tăng capacity so E05"),
        "E08": ("LGBM Poisson", "31 lá; lr=0,10; full", "Tăng learning rate so E05"),
        "E09": ("LGBM L1", "31 lá; lr=0,05; full", "Đổi objective so E05"),
        "E10": ("LGBM Poisson", "31 lá; lr=0,05; no promo/price", "Ablation khuyến mãi so E05"),
    }
    replacements["experiment_matrix"] = markdown_table(
        ["ID", "Thuật toán", "Cấu hình", "Mục tiêu"],
        [[s["id"], *descriptions[s["id"]]] for s in specs],
    )
    replacements["validation_results"] = markdown_table(
        ["ID", "WAPE CV mean±SD (%)", "WAPE pooled (%)", "Bias (%)", "MASE"],
        [
            [
                i,
                f"{r.wape_cv_mean:.3f} ± {r.wape_cv_sd:.3f}",
                f"{r.wape:.3f}",
                f"{r.bias:+.3f}",
                f"{r.mase_mean:.3f}",
            ]
            for i, r in results.iterrows()
        ],
    )
    replacements["validation_results"] += "\n\n" + markdown_table(
        ["ID", "Fit mean (s)", "Batch p95 mean (ms)", "Label coverage (%)", "Model gate"],
        [
            [
                i,
                f"{r.fit_seconds_mean:.4f}",
                f"{r.batch_p95_ms_mean:.3f}",
                f"{r.coverage * 100:.2f}",
                "Cơ sở" if i in ["E01", "E02", "E03"] else ("Đạt" if r.gate_pass else "Không đạt"),
            ]
            for i, r in results.iterrows()
        ],
    )
    delta_promo = results.loc["E10", "wape"] - results.loc["E05", "wape"]
    delta_leaves = results.loc["E07", "wape"] - results.loc["E05", "wape"]
    delta_lr = results.loc["E08", "wape"] - results.loc["E05", "wape"]
    improvement = selection["gates"][chosen]["relative_improvement"]
    replacements["results_analysis"] = (
        f"Baseline mạnh nhất trên validation là {baseline}, WAPE pooled={results.loc[baseline, 'wape']:.3f}%. "
        f"Cấu hình được khóa là {chosen}, WAPE pooled={results.loc[chosen, 'wape']:.3f}%, "
        f"cải thiện tương đối {improvement:.2f}% so với baseline đó. "
        f"Bias={results.loc[chosen, 'bias']:+.3f}% và worst group delta="
        f"{selection['gates'][chosen]['worst_group_delta_pp']:+.3f} điểm phần trăm; "
        f"trạng thái lựa chọn: {selection['status']}.\n\n"
        f"E10 bỏ promo/price có WAPE cao hơn E05 {delta_promo:.3f} điểm phần trăm, "
        "phù hợp với giả thuyết tín hiệu khuyến mãi hữu ích trong generator. "
        "Không diễn giải thành mức tăng doanh thu ở doanh nghiệp. "
        f"E07 tăng từ 31 lên 63 lá thay WAPE {delta_leaves:+.3f} điểm phần trăm; "
        f"fit time E05/E07 lần lượt {results.loc['E05', 'fit_seconds_mean']:.3f}/"
        f"{results.loc['E07', 'fit_seconds_mean']:.3f} giây. Capacity lớn hơn không mặc nhiên tốt hơn.\n\n"
        f"E08 tăng learning rate so E05 thay WAPE {delta_lr:+.3f} điểm phần trăm. "
        "Nếu chênh nhỏ so với biến động fold, xếp hạng theo quy tắc đã định trước chỉ là quyết định "
        "thực hành trong demo, không chứng minh ưu thế thống kê chắc chắn. "
        f"E04 Ridge đạt WAPE={results.loc['E04', 'wape']:.3f}% với fit mean "
        f"{results.loc['E04', 'fit_seconds_mean']:.3f} giây, là phương án đáng cân nhắc nếu "
        "chi phí/độ đơn giản quan trọng. Chưa có mô hình chi phí doanh nghiệp để kết luận "
        "lợi ích của chênh sai số đủ bù mọi chi phí vận hành.\n\n"
        f"Khuyến nghị nghiên cứu tiếp: giữ {chosen} làm candidate trong protocol này, "
        f"giữ {baseline} làm baseline khóa và fallback. Cấu hình đã chọn là "
        f"{json.dumps(selection['selected_spec'], ensure_ascii=False)}. "
        "Phương án production chỉ được chốt sau replay daily trên dữ liệu thực và ops/business gates."
    )
    replacements["test_results"] = (
        f"Final fit chỉ dùng nhãn đến {test['train_label_max']}; origin test từ "
        f"{test['test_origin_min']} đến {test['test_origin_max']}. "
        f"Có {test['eligible_test_rows']:,}/{test['all_test_rows']:,} dòng nhãn đủ điều kiện.\n\n"
        + markdown_table(
            ["Cấu hình khóa", "WAPE (%)", "Bias (%)", "MAE / RMSE", "MASE"],
            [
                [
                    i,
                    f"{s['wape']:.3f}",
                    f"{s['bias']:+.3f}",
                    f"{s['mae']:.3f} / {s['rmse']:.3f}",
                    f"{s['mase']:.3f}",
                ]
                for i, s in test["scores"].items()
            ],
        )
    )
    model_score, baseline_score = test["scores"][chosen], test["scores"][baseline]
    ci = test["uncertainty"]["difference_wape_pp_95ci"]
    replacements["test_analysis"] = (
        f"{chosen} cải thiện WAPE test tương đối {test['test_gate']['relative_improvement']:.2f}%; "
        f"test model gate {'đạt' if test['test_gate']['pass'] else 'không đạt'}. "
        f"Worst store/category delta là {test['test_gate']['worst_group_delta_pp']:+.3f} điểm phần trăm "
        "trên nhóm định trước có ít nhất 100 dòng và 1% lượng bán. "
        f"Baseline có bias={baseline_score['bias']:+.3f}%, model có bias={model_score['bias']:+.3f}%; "
        "cần kiểm tra bias cả ở từng nhóm, không chỉ tổng.\n\n"
        f"Chênh WAPE model−baseline={model_score['wape'] - baseline_score['wape']:+.3f} điểm phần trăm; "
        f"paired origin-block bootstrap 95%=[{ci[0]:+.3f}; {ci[1]:+.3f}], "
        f"{test['uncertainty']['blocks']} block, 2.000 lượt. "
        f"MASE tính được trên {model_score['mase_valid_series']} chuỗi, "
        f"loại {model_score['mase_excluded_series']} chuỗi vì scale/coverage. "
        f"Label coverage={model_score['coverage'] * 100:.2f}% chỉ là tỷ lệ nhãn không bị censor "
        "trong test, không phải ML serving coverage của production.\n\n"
        f"Final fit={model_score['fit_seconds']:.3f} giây; batch predict p95 nhỏ="
        f"{model_score['batch_p95_ms']:.3f} ms trên {model_score['n_test']:,} dòng. "
        "Không suy ra SLA trước 05:30 hoặc batch 280.000 dòng từ phép đo này."
    )
    replacements["registry_summary"] = (
        f"Đã tạo model registry thực tế: {registry['model_name']}, version {registry['version']}, "
        f"alias {registry['alias']}; run_id={registry['run_id']}. "
        "Registry cục bộ chỉ là candidate minh họa; production_promoted=false."
    )
    replacements["lineage_summary"] = (
        "Lineage thực nghiệm: raw SHA-256="
        + data["raw_sha256"]
        + "; feature SHA-256="
        + data["feature_sha256"]
        + ". Các hash đầy đủ cùng source hashes nằm trong manifests. "
        f"Model registry version {registry['version']} liên kết run {registry['run_id']}; "
        "selection.json ghi baseline/candidate, test_used_for_selection=false và gate reports. "
        "Manifest ghi chưa chạy Docker/Airflow; không có Git SHA phát hành được tự suy đoán."
    )
    checks = read_json("code_verification.json")
    replacements["verification_summary"] = markdown_table(
        ["Kiểm tra", "Kết quả thực hiện", "Ý nghĩa"],
        [
            ["Pytest", checks["pytest_summary"], "Contracts/leakage/metric/gate tests"],
            ["Ruff", checks["ruff_summary"], "Lỗi tên/import và format"],
            ["Python compile", checks["compile_summary"], "Cú pháp src, CLI, DAG và exporter"],
            [
                "Experiment/MLflow",
                "30 fold runs thành công + final model run",
                "Bằng chứng CSV/JSON/registry cục bộ",
            ],
            [
                "Airflow/Docker/ERP",
                "Chưa chạy tích hợp",
                "Thiết kế và code sample; cần kiểm chứng staging",
            ],
        ],
    )
    replacements["verification_summary"] += (
        f"\n\nMôi trường thực thi: Python {environment['python']}, "
        f"LightGBM {environment['packages']['lightgbm']}, MLflow {environment['packages']['mlflow']}, "
        f"scikit-learn {environment['packages']['scikit-learn']}; {environment['platform']}. "
        "Lượt kỹ thuật ban đầu được thay thế sau khi sửa purge toàn bộ origin, cutoff theo tháng "
        "và tương thích serialization. Không thay ma trận, seed hoặc ngưỡng để cải thiện test. "
        "Chỉ bảng/artifacts của lượt hoàn tất với source hash hiện tại được dùng trong báo cáo."
    )
    code_blocks = {
        "config_json": (ROOT / "configs/demo.json").read_text(encoding="utf-8"),
        "config_loader": """import os
from pathlib import Path
from src.forecasting import Config, load_config

# load_config parses JSON, validates types/gates, then applies documented overrides.
os.environ["DDM_SEED"] = "501"
config: Config = load_config(Path("configs/demo.json"))""",
        "seeds": function_source("src/forecasting.py", "set_seeds")
        + "\n\n# Also passed to LGBMRegressor:\nrandom_state = 501\ndeterministic = True\nforce_col_wise = True",
        "dockerfile": (ROOT / "Dockerfile").read_text(encoding="utf-8"),
    }
    code_blocks["airflow_dag"] = '''from datetime import datetime, timedelta, timezone
from airflow.sdk import dag, task
from retail_demo import run_stage, failure_alert

@dag(schedule="0 18 * * 6", catchup=False, max_active_runs=1,
     start_date=datetime(2026, 10, 3, tzinfo=timezone.utc),
     default_args={"retries": 2, "retry_delay": timedelta(minutes=5),
                   "retry_exponential_backoff": True,
                   "on_failure_callback": failure_alert})
def candidate_workflow():
    """Weekly demo candidate; tasks exchange artifact paths only."""
    @task(retries=0)
    def prepare(output: str) -> str:
        return run_stage("prepare", output)
    @task
    def train_validate(uri: str) -> str:
        return run_stage("search", uri)
    @task(retries=0)
    def evaluate(uri: str) -> str:
        return run_stage("evaluate", uri)
    @task
    def register(uri: str) -> str:
        return run_stage("register", uri)
    register(evaluate(train_validate(prepare("/app/artifacts/demo"))))

demo_dag = candidate_workflow()
# Compact illustration; full DAG adds data-interval paths and execution_timeout.'''
    code_blocks["mlflow_logging"] = """mlflow.set_tracking_uri(config.tracking_uri)
mlflow.set_experiment(config.experiment_name)
with mlflow.start_run(run_name=f"{spec['id']}-fold{fold}") as run:
    mlflow.set_tags({"evidence_type": "synthetic", "phase": "validation",
                     "data_sha256": manifest["raw_sha256"]})
    mlflow.log_params({**spec, "seed": config.seed, "n_jobs": config.n_jobs,
                       "n_estimators": config.n_estimators, "fold": fold,
                       "validation_start": start, "validation_end": end})
    model.fit(train[features], train.y.to_numpy())
    predictions, latency = benchmark(model, valid[features])
    score = metrics(valid.y.to_numpy(), predictions)
    mlflow.log_metrics({k: float(v) for k, v in score.items() if v is not None})
    mlflow.log_dict({"features": features, "spec": spec}, "feature_config.json")"""
    code_blocks[
        "model_registration"
    ] = """# Sau khi model đã fit và selection đã khóa (run_experiments.py).
signature = infer_signature(example, np.maximum(0.0, model.predict(example)))
with mlflow.start_run(run_name="locked-test-demo-candidate") as run:
    mlflow.set_tags({"evidence_type": "synthetic", "deployment_allowed": "false"})
    mlflow.log_dict(asdict(config), "config.json")
    mlflow.log_artifact(str(out / "selection.json"))
    mlflow.log_artifact(str(out / "test_results.json"))
    info = mlflow.sklearn.log_model(
        model, name="model", signature=signature, input_example=example,
        skops_trusted_types=["collections.OrderedDict", "lightgbm.basic.Booster",
                             "lightgbm.sklearn.LGBMRegressor", "numpy.dtype"],
        registered_model_name=config.registered_model_name,
    )
    version = str(info.registered_model_version)
    client = MlflowClient()
    client.set_model_version_tag(config.registered_model_name, version,
                                 "evidence_type", "synthetic")
    client.set_model_version_tag(config.registered_model_name, version,
                                 "deployment_allowed", "false")
    client.set_registered_model_alias(config.registered_model_name,
                                      "demo_candidate", version)
# Serving clips output >= 0 and pins the numeric version; no champion promotion."""
    for label, block in code_blocks.items():
        replacements["code:" + label] = (
            "```"
            + ("json" if label == "config_json" else "python")
            + "\n"
            + block.rstrip()
            + "\n```"
        )
    return replacements


class Diagram(Flowable):
    """Vector diagrams with readable Vietnamese labels and explicit branch arrows."""

    def __init__(self, kind: str):
        super().__init__()
        self.kind = kind
        self.width = WIDTH
        self.height = 365 if kind == "pipeline" else 345

    def draw(self) -> None:
        c = self.canv
        color = colors.HexColor

        def box(x, y, w, h, lines, fill="#e5f0f3"):
            c.setFillColor(color(fill))
            c.setStrokeColor(color("#7395a3"))
            c.roundRect(x, y, w, h, 5, fill=1, stroke=1)
            c.setFillColor(color("#19384d"))
            c.setFont("ArialVN", 8)
            for i, label in enumerate(lines):
                c.drawCentredString(x + w / 2, y + h - 13 - i * 11, label)

        def arrow(points, dashed=False):
            c.setStrokeColor(color("#476978"))
            c.setLineWidth(1)
            c.setDash(3, 2) if dashed else c.setDash()
            for first, second in zip(points, points[1:]):
                c.line(*first, *second)
            x1, y1 = points[-2]
            x2, y2 = points[-1]
            angle = math.atan2(y2 - y1, x2 - x1)
            for offset in [-0.45, 0.45]:
                c.line(x2, y2, x2 - 5 * math.cos(angle + offset), y2 - 5 * math.sin(angle + offset))
            c.setDash()

        if self.kind == "pipeline":
            box(0, 312, 140, 45, ["POS / tồn kho / plans", "Cutoff 01:00, o=D−1"])
            box(174, 312, 140, 45, ["S1 Ingestion", "Raw bất biến + manifest"])
            box(348, 312, 139, 45, ["S2 Validation", "Schema / keys / coverage"])
            arrow([(140, 334), (174, 334)])
            arrow([(314, 334), (348, 334)])
            box(
                348,
                229,
                139,
                49,
                ["S3 Split + label mask", "S4 Preprocess (fit train)", "S5 Point-in-time features"],
            )
            arrow([(418, 312), (418, 278)])
            box(
                174,
                229,
                140,
                49,
                ["S6 Training / CV", "Baseline + 10 cấu hình", "MLflow runs / artifacts"],
            )
            arrow([(348, 254), (314, 254)])
            box(
                0,
                229,
                140,
                49,
                ["S7 Evaluate + gates", "Selection → locked test", "Data/model scorecard"],
            )
            arrow([(174, 254), (140, 254)])
            box(
                0,
                143,
                140,
                49,
                ["S8 Registry → staging", "Shadow / ops / phê duyệt", "Pin version + rollback"],
            )
            arrow([(70, 229), (70, 192)])
            box(
                174,
                143,
                140,
                49,
                ["S9 Inference hằng ngày", "Snapshot + pinned model", "Forecast 1–7 / fallback"],
            )
            arrow([(140, 167), (174, 167)])
            arrow([(418, 229), (418, 207), (244, 207), (244, 192)])
            box(
                348,
                143,
                139,
                49,
                ["Forecast + draft orders", "Nonnegative / đủ h", "IP, ETA, pack size"],
            )
            arrow([(314, 167), (348, 167)])
            box(
                348,
                58,
                139,
                49,
                ["S10 Atomic publish", "PostgreSQL → API", "Nhân viên duyệt → ERP"],
            )
            arrow([(418, 143), (418, 107)])
            box(
                174,
                58,
                140,
                49,
                ["Monitoring + labels", "SLO / coverage / drift", "Metrics khi nhãn đủ tuổi"],
                "#edf2e2",
            )
            arrow([(348, 83), (314, 83)])
            arrow([(244, 107), (244, 127), (330, 127), (330, 298), (244, 298), (244, 278)], True)
            box(
                0,
                58,
                140,
                49,
                ["Quarantine / alert", "Dữ liệu sai → cách ly", "Không tự điền sales=0"],
                "#f9ece8",
            )
            arrow([(418, 312), (418, 294), (155, 294), (155, 83), (140, 83)], True)
            c.setFont("ArialVN-Italic", 7.4)
            c.drawString(
                0, 25, "Nét liền: dữ liệu/phụ thuộc chính. Nét đứt: lỗi hoặc vòng phản hồi."
            )
            c.drawString(
                0,
                12,
                "Huấn luyện tuần tách khỏi inference ngày; model gate không thay ops/business approval.",
            )
        else:
            c.setFont("ArialVN-Bold", 9)
            c.drawString(0, 330, "DAG NGÀY — kế hoạch vận hành; chưa tích hợp POS/ERP")
            box(0, 256, 105, 49, ["wait/ingest", "raw manifest"])
            box(126, 256, 105, 49, ["validate", "clean/quarantine"])
            box(252, 256, 105, 49, ["preprocess", "features @ origin"])
            box(378, 256, 109, 49, ["predict partitions", "pin model version"])
            for end, start in [(105, 126), (231, 252), (357, 378)]:
                arrow([(end, 280), (start, 280)])
            box(378, 179, 109, 49, ["forecast gate", "draft order rules"])
            box(252, 179, 105, 49, ["publish atomic", "05:30 deadline"])
            box(126, 179, 105, 49, ["monitor / alert", "technical status"], "#edf2e2")
            arrow([(433, 256), (433, 228)])
            arrow([(378, 204), (357, 204)])
            arrow([(252, 204), (231, 204)])
            c.drawString(
                0, 149, "DAG TUẦN — ứng viên; release cần dữ liệu thực + ops gate + người duyệt"
            )
            box(0, 79, 105, 49, ["snapshot / split", "label cutoff"])
            box(126, 79, 105, 49, ["train / CV", "selection locked"])
            box(252, 79, 105, 49, ["locked evaluation", "model scorecard"])
            box(378, 79, 109, 49, ["stage / review", "registry candidate"])
            for end, start in [(105, 126), (231, 252), (357, 378)]:
                arrow([(end, 103), (start, 103)])
            c.setFont("ArialVN", 8)
            c.drawString(
                0, 43, "DAG demo trong mã: prepare → search → evaluate → register (synthetic only)."
            )
            c.drawString(
                0, 28, "prepare gộp ingest/validate/features; register không phát hành production."
            )
            c.drawString(
                0,
                13,
                "XCom chỉ mang URI/metadata; task nặng chạy CLI; artifact storage dùng chung.",
            )


def save_diagrams() -> None:
    """Export diagrams as standalone PDF/PNG and Mermaid source."""
    for kind in ["pipeline", "dag"]:
        diagram = Diagram(kind)
        pdf = OUT / (kind + ".pdf")
        drawing = canvas.Canvas(str(pdf), pagesize=(WIDTH + 20, diagram.height + 20))
        diagram.drawOn(drawing, 10, 10)
        drawing.save()
        with pymupdf.open(pdf) as document:
            document[0].get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(str(OUT / (kind + ".png")))
    (OUT / "pipeline.mmd").write_text(
        """flowchart LR
  source[POS / Inventory / Published plans] --> ingest[Ingest / Raw manifest]
  ingest --> validate[Validation]
  validate --> features[Purged split / Train-fitted preprocessing / Point-in-time features]
  validate -. fail .-> quarantine[Quarantine / Alert]
  features --> train[10 configurations / temporal CV]
  train --> evaluate[Lock selection / Test / Model gates]
  evaluate --> registry[Registry / Staging / Ops gates / Human review]
  registry --> predict[Daily batch with pinned model]
  features --> predict
  predict --> order[Forecast gate / Draft orders]
  order --> publish[Atomic publish / API / Approval / ERP]
  publish --> monitor[Telemetry / Matured labels]
  monitor -. investigation .-> train
""",
        encoding="utf-8",
    )
    (OUT / "dag.mmd").write_text(
        """flowchart LR
  subgraph Daily[Daily operational design]
    ingest --> validate --> features --> predict --> forecast_gate --> draft_orders --> publish
    publish --> technical_monitor
  end
  subgraph Weekly[Weekly candidate design]
    snapshot --> split --> train_cv --> select --> locked_evaluate --> stage_review
  end
  subgraph Demo[Implemented demo DAG definition]
    prepare_demo --> search_demo --> evaluate_demo --> register_demo_only
  end
""",
        encoding="utf-8",
    )


class ReportDocument(SimpleDocTemplate):
    """Add bookmarks and a live PDF table of contents."""

    def afterFlowable(self, flowable):
        if hasattr(flowable, "report_heading"):
            label, level, key = flowable.report_heading
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(label, key, level=level, closed=False)
            if label != "Mục lục":
                self.notify("TOCEntry", (level, label, self.page, key))


def inline(value: str) -> str:
    """Escape text, preserve simple Markdown emphasis and click-through URLs."""
    value = escape(value)
    value = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"`([^`]+)`", r'<font name="CodeVN" size="9">\1</font>', value)
    value = re.sub(r"(https?://[^\s<>]+)", r'<link href="\1" color="#146b73">\1</link>', value)
    # Break long provenance hashes safely, preserving their characters in extracted text.
    value = re.sub(r"\b[0-9a-f]{64}\b", lambda m: "<br/>" + m.group(0), value)
    return value


def export_report(source: str) -> None:
    """Parse report Markdown into PDF and editable Word, including code and figures."""
    story = []
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = Pt(A4[0]), Pt(A4[1])
    section.left_margin = section.right_margin = Pt(54)
    section.top_margin, section.bottom_margin = Pt(48), Pt(54)
    section.different_first_page_header_footer = True
    normal = word.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(8)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        word.styles[name].font.name = "Times New Roman"
        word.styles[name].paragraph_format.keep_with_next = True
    section.header.paragraphs[0].text = "DDM501 | Assignment 2 | ML Pipeline & MLOps"
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    section.footer.paragraphs[0]._p.append(page_field)
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    word.settings.element.append(update)
    count_heading, count_table, count_figure = 0, 0, 0

    def heading(label, level):
        nonlocal count_heading
        count_heading += 1
        paragraph = Paragraph(inline(label), styles[f"Heading{level}"])
        paragraph.report_heading = (label, level - 1, f"heading-{count_heading}")
        story.append(paragraph)
        word.add_heading(label, level=level)

    def text(label):
        story.append(Paragraph(inline(label), styles["BodyVN"]))
        paragraph = word.add_paragraph(label)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def table(lines):
        nonlocal count_table
        count_table += 1
        rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
        rows = rows[:1] + rows[2:]
        columns = len(rows[0])
        caption = f"Bảng {count_table}. {rows[0][0]} và các tiêu chí đối chiếu."
        story.append(Paragraph(escape(caption), styles["CaptionVN"]))
        word.add_paragraph(caption)
        widths = [WIDTH / columns] * columns
        if columns == 2:
            widths = [WIDTH * 0.30, WIDTH * 0.70]
        elif columns == 3:
            widths = [WIDTH * 0.20, WIDTH * 0.37, WIDTH * 0.43]
        elif columns == 4:
            widths = [WIDTH * 0.09, WIDTH * 0.23, WIDTH * 0.35, WIDTH * 0.33]
        elif columns == 5:
            widths = [WIDTH * 0.15, WIDTH * 0.26, WIDTH * 0.21, WIDTH * 0.19, WIDTH * 0.19]
        cells = [[Paragraph(inline(cell), styles["CellVN"]) for cell in row] for row in rows]
        tab = Table(cells, colWidths=widths, repeatRows=1, hAlign="LEFT")
        tab.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dcecf1")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor("#8fb1bd")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#f4f7fa")],
                    ),
                    ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor("#dfe5eb")),
                ]
            )
        )
        tab.spaceAfter = 10
        story.append(tab)
        word_table = word.add_table(rows=1, cols=columns)
        word_table.style = "Light Shading Accent 1"
        for cell, value in zip(word_table.rows[0].cells, rows[0]):
            cell.text = value
        for row in rows[1:]:
            for cell, value in zip(word_table.add_row().cells, row):
                cell.text = value
        for row in word_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9.5)
        repeat = OxmlElement("w:tblHeader")
        word_table.rows[0]._tr.get_or_add_trPr().append(repeat)

    lines = source.splitlines()
    divider = lines.index("---")
    story.append(Spacer(1, 65))
    for line in lines[:divider]:
        if not line.strip():
            continue
        label = re.sub(r"^#+\s*", "", line)
        story.append(
            Paragraph(escape(label), styles["Cover"] if line.startswith("#") else styles["Meta"])
        )
        paragraph = word.add_paragraph(label)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if line.startswith("#"):
            paragraph.runs[0].bold = True
            paragraph.runs[0].font.size = Pt(22)
    story.append(PageBreak())
    word.add_page_break()
    heading("Mục lục", 1)
    toc = TableOfContents()
    toc.levelStyles = [
        ParagraphStyle("TOC1", fontName="TimesVN-Bold", fontSize=10.5, leading=15, spaceBefore=5),
        ParagraphStyle(
            "TOC2", fontName="TimesVN", fontSize=10, leading=13.5, leftIndent=14, spaceBefore=1
        ),
    ]
    story.append(toc)
    toc_field = OxmlElement("w:fldSimple")
    toc_field.set(qn("w:instr"), 'TOC \\o "1-2" \\h \\z \\u')
    word.add_paragraph()._p.append(toc_field)
    story.append(PageBreak())
    word.add_page_break()
    i = divider + 1
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "---":
            story.append(PageBreak())
            word.add_page_break()
        elif line.startswith("# "):
            heading(line[2:], 1)
        elif line.startswith("## "):
            heading(line[3:], 2)
        elif line.startswith("### "):
            heading(line[4:], 3)
        elif line.startswith("{{figure:"):
            nonlocal_kind = line[len("{{figure:") : -2]
            count_figure += 1
            if nonlocal_kind in ["pipeline", "dag"]:
                story.append(Diagram(nonlocal_kind))
                filename = OUT / (nonlocal_kind + ".png")
            else:
                filename = ART / (nonlocal_kind + ".png")
                with pymupdf.open(filename) as photo:
                    ratio = photo[0].rect.height / photo[0].rect.width
                story.append(Image(str(filename), width=WIDTH, height=WIDTH * ratio))
            word.add_picture(str(filename), width=Inches(WIDTH / 72))
            label = {
                "pipeline": "Kiến trúc pipeline đề xuất và điểm kiểm soát",
                "dag": "DAG vận hành ngày/tuần và phạm vi demo",
                "validation_comparison": "WAPE ba fold xác thực trên dữ liệu mô phỏng",
                "test_horizons": "WAPE theo chân trời của candidate và baseline đã khóa",
            }[nonlocal_kind]
            caption = f"Hình {count_figure}. {label}. Nguồn: thiết kế hoặc artifacts thực nghiệm của tác giả."
            story.append(Paragraph(escape(caption), styles["CaptionVN"]))
            word.add_paragraph(caption)
        elif line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            code = "\n".join(block)
            longest = max(
                (pdfmetrics.stringWidth(value, "CodeVN", 7.1) for value in block), default=0
            )
            code_style = ParagraphStyle(f"CodeBlock{count_heading}-{i}", parent=styles["CodeVN"])
            if longest > WIDTH:
                code_style.fontSize = max(5.4, 7.1 * WIDTH / longest)
                code_style.leading = code_style.fontSize * 1.32
            story.append(Preformatted(code, code_style))
            paragraph = word.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(code)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        elif line.startswith("|"):
            tab = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tab.append(lines[i])
                i += 1
            table(tab)
            continue
        else:
            pieces = [line]
            while (
                i + 1 < len(lines)
                and lines[i + 1].strip()
                and not lines[i + 1].startswith(("#", "|", "```", "{{figure:"))
                and lines[i + 1].strip() != "---"
            ):
                i += 1
                pieces.append(lines[i].strip())
            text(" ".join(pieces))
        i += 1

    def footer(c, d):
        if d.page == 1:
            return
        c.saveState()
        c.setStrokeColor(colors.HexColor("#cbd9e1"))
        c.line(54, 43, A4[0] - 54, 43)
        c.setFillColor(colors.HexColor("#536b80"))
        c.setFont("ArialVN", 8)
        c.drawString(54, 30, "DDM501 | Assignment 2 | Nguyễn Minh Trường — 25MS13312")
        c.drawRightString(A4[0] - 54, 30, str(d.page))
        c.restoreState()

    pdf = OUT / (STEM + ".pdf")
    ReportDocument(
        str(pdf),
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=48,
        bottomMargin=58,
        title="DDM501 Assignment 2 — Retail ML Pipeline & MLOps",
        author="Nguyễn Minh Trường",
    ).multiBuild(story, onFirstPage=footer, onLaterPages=footer)
    word.save(OUT / (STEM + ".docx"))
    published = source
    for figure in ["pipeline", "dag", "validation_comparison", "test_horizons"]:
        published = published.replace("{{figure:" + figure + "}}", f"![{figure}]({figure}.png)")
        if figure not in ["pipeline", "dag"]:
            (OUT / (figure + ".png")).write_bytes((ART / (figure + ".png")).read_bytes())
    (OUT / "REPORT.md").write_text(published, encoding="utf-8")
    (OUT / "build_manifest.json").write_text(
        json.dumps(
            {
                "pdf": pdf.name,
                "tables": count_table,
                "figures": count_figure,
                "headings": count_heading,
                "report_source_sha256": hashlib.sha256(source.encode()).hexdigest(),
            },
            indent=2,
        ),
        encoding="utf-8",
    )
    print(
        f"Report generated; {count_table} tables, {count_figure} figures, {count_heading} headings"
    )


def main() -> None:
    """Resolve evidence sections before rendering; reject leftover placeholders."""
    save_diagrams()
    source = (ROOT / "REPORT_TEMPLATE_VI.md").read_text(encoding="utf-8")
    for key, value in generated_sections().items():
        source = source.replace("{{" + key + "}}", value)
    unresolved = re.findall(r"\{\{(?!figure:)[^}]+\}\}", source)
    if unresolved:
        raise ValueError(f"Unresolved report placeholders: {unresolved}")
    (ROOT / "REPORT.md").write_text(source, encoding="utf-8")
    export_report(source)


if __name__ == "__main__":
    main()
